import docx
import os

doc = docx.Document('docs/Copy of UrbanFix_AI_Final_Report.docx')
print(f'Total paragraphs: {len(doc.paragraphs)}')
print(f'Total tables: {len(doc.tables)}')
print(f'Total inline shapes: {len(doc.inline_shapes)}')

for i, p in enumerate(doc.paragraphs):
    t = p.text.strip()
    if t.startswith('Chapter') or t.startswith('Table') or t.startswith('Figure') or t.startswith('1.') or t.startswith('2.') or t.startswith('3.') or t.startswith('4.') or t.startswith('5.') or t.startswith('6.'):
        print(f'Para {i}: {t[:100]}')
